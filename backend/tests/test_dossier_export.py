"""Export PDF du dossier construit par la chaine.

Ce qui compte ici n'est pas que le PDF existe, mais qu'il ne mente jamais sur
son etat : un dossier non valide ne doit pas pouvoir passer pour un document
final, et les constats internes ne doivent pas se retrouver dans ce qu'on
envoie a un financeur.
"""

from __future__ import annotations

import json

import pytest

from app.services.ai.base import AICompletionResponse, AIProvider
from tests.conftest import job_result

PATCHES = {
    "DEVELOPMENT": {"public cible": "18-35, urbain"},
    "PRODUCER": {"jours de tournage": 24},
}


class Scripted(AIProvider):
    name = "scripted"

    def __init__(self, verdict: str = "PASS", findings: list | None = None) -> None:
        self.verdict = verdict
        self.findings = findings or []

    def complete(self, request):
        role = request.metadata["agent_role"]
        patch = PATCHES.get(role)
        return AICompletionResponse(
            text=json.dumps(
                {
                    "analysis": f"## {role}",
                    "rationale": "…",
                    "decisions": [],
                    "modifications": {
                        "preserved": [],
                        "modified": [],
                        "removed": [],
                        "added": [],
                        "reasoning": [],
                    },
                    "findings": self.findings if role.endswith("VALIDATOR") else [],
                    "verdict": self.verdict if role.endswith("VALIDATOR") else None,
                    "state_patch": patch,
                    "next_agent_instructions": "",
                },
                ensure_ascii=False,
            ),
            model="m",
            provider=self.name,
        )


def _blocking():
    return [
        {
            "severity": "CRITICAL",
            "element": "budget",
            "description": "Aucun budget fourni.",
            "owner": "PRODUCER",
            "suggested_correction": "Chiffrer les 24 jours de tournage.",
        }
    ]


@pytest.fixture
def paid(client, make_user):
    """L'export est réservé aux offres qui l'incluent."""
    headers, _ = make_user("export@example.com", plan="PRO_AUTHOR")
    project_id = client.post(
        "/api/v1/projects",
        json={"title": "Taxi 237", "project_type": "FEATURE_FILM"},
        headers=headers,
    ).json()["id"]
    return headers, project_id


def _run_chain(client, headers, project_id, monkeypatch, provider):
    monkeypatch.setattr("app.services.ai.service.build_provider", lambda *a, **k: provider)
    job_result(client.post(f"/api/v1/projects/{project_id}/dossier/run", headers=headers))


def _export(client, headers, project_id):
    return client.get(f"/api/v1/projects/{project_id}/export/dossier/pdf", headers=headers)


# ----------------------------------------------------------------------


def test_nothing_to_export_before_the_chain_has_run(client, paid):
    """Un PDF vide serait plus déroutant qu'une erreur qui dit quoi faire."""
    headers, project_id = paid
    response = _export(client, headers, project_id)
    assert response.status_code == 400
    assert response.json()["code"] == "dossier_not_built"


def test_a_validated_dossier_exports_as_a_pdf(client, paid, monkeypatch):
    headers, project_id = paid
    _run_chain(client, headers, project_id, monkeypatch, Scripted())

    response = _export(client, headers, project_id)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    # Le nom du fichier ne porte pas la mention brouillon.
    assert "BROUILLON" not in response.headers["content-disposition"]


def test_an_unvalidated_dossier_is_named_a_draft(client, paid, monkeypatch):
    """Le nom du fichier doit dire l'état : il survit au téléchargement."""
    headers, project_id = paid
    _run_chain(
        client, headers, project_id, monkeypatch, Scripted("BLOCKED", _blocking())
    )

    response = _export(client, headers, project_id)
    assert response.status_code == 200
    assert "BROUILLON" in response.headers["content-disposition"]


def test_the_draft_says_so_inside_the_document(client, paid, monkeypatch):
    """Un fichier renommé perdrait l'avertissement : il est aussi dans la page."""
    from pypdf import PdfReader

    headers, project_id = paid
    _run_chain(
        client, headers, project_id, monkeypatch, Scripted("BLOCKED", _blocking())
    )

    text = _pdf_text(PdfReader, _export(client, headers, project_id).content)
    assert "BROUILLON" in text
    assert "ne doit pas être soumis" in text
    assert "bloqué" in text


def test_a_validated_dossier_never_claims_to_be_a_guarantee(client, paid, monkeypatch):
    """Règle produit : le score et le dossier aident, ils ne garantissent rien."""
    from pypdf import PdfReader

    headers, project_id = paid
    _run_chain(client, headers, project_id, monkeypatch, Scripted())

    text = _pdf_text(PdfReader, _export(client, headers, project_id).content)
    assert "relecture humaine" in text
    assert "BROUILLON" not in text


def test_internal_findings_stay_out_of_a_validated_dossier(client, paid, monkeypatch):
    """Les exposer à un financeur desservirait le projet.

    Un constat mineur ne bloque pas l'export ; il ne doit pas pour autant se
    retrouver dans le document qu'un comité va lire.
    """
    from pypdf import PdfReader

    headers, project_id = paid
    minor = [
        {
            "severity": "MINOR",
            "element": "titre",
            "description": "Formulation perfectible du titre.",
            "owner": "DEVELOPMENT",
        }
    ]
    _run_chain(
        client, headers, project_id, monkeypatch, Scripted("PASS_WITH_WARNINGS", minor)
    )

    response = _export(client, headers, project_id)
    assert "BROUILLON" not in response.headers["content-disposition"]

    text = _pdf_text(PdfReader, response.content)
    assert "Points à traiter" not in text
    assert "Formulation perfectible" not in text


def test_a_draft_carries_its_findings_and_their_owner(client, paid, monkeypatch):
    """C'est la seule raison de produire ce PDF-là : savoir quoi corriger."""
    from pypdf import PdfReader

    headers, project_id = paid
    _run_chain(
        client, headers, project_id, monkeypatch, Scripted("BLOCKED", _blocking())
    )

    text = _pdf_text(PdfReader, _export(client, headers, project_id).content)
    assert "Points à traiter" in text
    assert "Aucun budget fourni" in text
    assert "Production" in text  # l'agent capable de corriger
    assert "Chiffrer les 24 jours" in text


def test_the_sections_produced_by_the_agents_are_in_the_pdf(client, paid, monkeypatch):
    from pypdf import PdfReader

    headers, project_id = paid
    _run_chain(client, headers, project_id, monkeypatch, Scripted())

    text = _pdf_text(PdfReader, _export(client, headers, project_id).content)
    assert "Concept" in text
    assert "18-35, urbain" in text
    assert "Plan de production" in text


def test_an_empty_section_is_not_printed(client, paid, monkeypatch):
    """Un titre suivi de rien se lit comme un défaut de mise en page."""
    from pypdf import PdfReader

    headers, project_id = paid
    _run_chain(client, headers, project_id, monkeypatch, Scripted())

    text = _pdf_text(PdfReader, _export(client, headers, project_id).content)
    # Aucun agent n'a rempli l'analyse culturelle dans ce scénario.
    assert "Analyse culturelle" not in text


def test_another_user_cannot_export_this_dossier(client, paid, make_user, monkeypatch):
    headers, project_id = paid
    _run_chain(client, headers, project_id, monkeypatch, Scripted())

    intruder, _ = make_user("intrus-export@example.com", plan="PRO_AUTHOR")
    assert _export(client, intruder, project_id).status_code == 404


def test_the_free_plan_cannot_export(client, make_user):
    """L'export est une fonctionnalité d'offre, dossier compris."""
    headers, _ = make_user("gratuit-export@example.com")
    project_id = client.post(
        "/api/v1/projects",
        json={"title": "Projet gratuit", "project_type": "DOCUMENTARY"},
        headers=headers,
    ).json()["id"]
    assert _export(client, headers, project_id).status_code == 402


def _pdf_text(reader_cls, content: bytes) -> str:
    import io

    reader = reader_cls(io.BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


# ----------------------------------------------------------------------
# Word : mêmes règles, format modifiable


def _docx(client, headers, project_id):
    return client.get(f"/api/v1/projects/{project_id}/export/dossier/docx", headers=headers)


def _docx_text(content: bytes) -> str:
    import io

    from docx import Document as DocxDocument

    return "\n".join(p.text for p in DocxDocument(io.BytesIO(content)).paragraphs)


def test_nothing_to_export_in_word_either(client, paid):
    headers, project_id = paid
    assert _docx(client, headers, project_id).json()["code"] == "dossier_not_built"


def test_a_validated_dossier_exports_as_a_word_file(client, paid, monkeypatch):
    headers, project_id = paid
    _run_chain(client, headers, project_id, monkeypatch, Scripted())

    response = _docx(client, headers, project_id)
    assert response.status_code == 200
    assert "wordprocessingml" in response.headers["content-type"]
    # Un .docx est une archive zip : la signature le prouve mieux qu'un nom.
    assert response.content.startswith(b"PK")
    assert "BROUILLON" not in response.headers["content-disposition"]


def test_the_word_draft_carries_the_same_warning(client, paid, monkeypatch):
    headers, project_id = paid
    _run_chain(client, headers, project_id, monkeypatch, Scripted("BLOCKED", _blocking()))

    response = _docx(client, headers, project_id)
    assert "BROUILLON" in response.headers["content-disposition"]

    text = _docx_text(response.content)
    assert "BROUILLON" in text
    assert "ne doit pas être soumis" in text
    assert "bloqué" in text


def test_the_two_formats_say_the_same_thing(client, paid, monkeypatch):
    """Si les avertissements divergeaient, un format serait moins clair.

    C'est la raison d'être des textes partagés : ce test échoue dès que l'un
    des deux exports se met à raconter autre chose.
    """
    from pypdf import PdfReader

    headers, project_id = paid
    _run_chain(client, headers, project_id, monkeypatch, Scripted("BLOCKED", _blocking()))

    pdf = _pdf_text(PdfReader, _export(client, headers, project_id).content)
    word = _docx_text(_docx(client, headers, project_id).content)

    for phrase in (
        "BROUILLON — dossier non validé.",
        "ne doit pas être soumis en l'état",
        "Points à traiter",
        "Aucun budget fourni",
        "Chiffrer les 24 jours",
    ):
        assert phrase in pdf, f"absent du PDF : {phrase}"
        assert phrase in word, f"absent du Word : {phrase}"


def test_word_findings_stay_out_of_a_validated_dossier(client, paid, monkeypatch):
    headers, project_id = paid
    minor = [
        {
            "severity": "MINOR",
            "element": "titre",
            "description": "Formulation perfectible du titre.",
            "owner": "DEVELOPMENT",
        }
    ]
    _run_chain(client, headers, project_id, monkeypatch, Scripted("PASS_WITH_WARNINGS", minor))

    text = _docx_text(_docx(client, headers, project_id).content)
    assert "Points à traiter" not in text
    assert "Formulation perfectible" not in text
    assert "relecture humaine" in text


def test_the_word_document_stays_editable(client, paid, monkeypatch):
    """Tout l'intérêt du format : un financeur qui annote doit pouvoir écrire."""
    import io

    from docx import Document as DocxDocument

    headers, project_id = paid
    _run_chain(client, headers, project_id, monkeypatch, Scripted())

    document = DocxDocument(io.BytesIO(_docx(client, headers, project_id).content))
    document.add_paragraph("Annotation d'un lecteur.")
    again = io.BytesIO()
    document.save(again)
    assert "Annotation d'un lecteur." in _docx_text(again.getvalue())


def test_the_free_plan_cannot_export_word_either(client, make_user):
    headers, _ = make_user("gratuit-docx@example.com")
    project_id = client.post(
        "/api/v1/projects",
        json={"title": "Projet gratuit", "project_type": "DOCUMENTARY"},
        headers=headers,
    ).json()["id"]
    assert _docx(client, headers, project_id).status_code == 402


def test_the_word_draft_warning_is_visually_distinct(client, paid, monkeypatch):
    """Le gras seul se confond avec un intertitre : c'est la couleur qui alerte.

    Testé parce qu'une régression ici est invisible à la relecture du code et
    ne casse aucune autre assertion — le document resterait « correct ».
    """
    import io

    from docx import Document as DocxDocument

    headers, project_id = paid
    _run_chain(client, headers, project_id, monkeypatch, Scripted("BLOCKED", _blocking()))

    document = DocxDocument(io.BytesIO(_docx(client, headers, project_id).content))
    lead = next(
        run
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if "BROUILLON" in run.text
    )
    assert lead.bold
    assert lead.font.color.rgb is not None, "un avertissement en noir passe inaperçu"


def test_a_validated_word_dossier_has_no_red_warning(client, paid, monkeypatch):
    import io

    from docx import Document as DocxDocument

    headers, project_id = paid
    _run_chain(client, headers, project_id, monkeypatch, Scripted())

    document = DocxDocument(io.BytesIO(_docx(client, headers, project_id).content))
    lead = next(
        run
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if "Dossier contrôlé" in run.text
    )
    assert lead.bold
    assert lead.font.color.rgb is None


# ----------------------------------------------------------------------
# Un contrôle interrompu ne se lit pas comme un contrôle sévère


def _notice(exportable: bool, interrupted: bool) -> tuple[str, str]:
    from app.services.export_service import ExportService

    return ExportService._dossier_notice(exportable, interrupted)


def test_an_interrupted_run_says_the_control_did_not_happen():
    """Le défaut corrigé : « les contrôles ont relevé des points à traiter ».

    Pour une chaîne arrêtée en chemin, c'est faux deux fois — les contrôles
    n'ont rien relevé, et ils n'ont pas tourné. Le lecteur en conclurait que
    le dossier a été lu et jugé perfectible, alors qu'il n'a pas été lu.
    """
    from app.services.export_service import DRAFT_WARNING, INTERRUPTED_WARNING

    lead, rest = _notice(exportable=False, interrupted=True)
    assert (lead, rest) == INTERRUPTED_WARNING
    assert (lead, rest) != DRAFT_WARNING
    assert "interrompu" in lead.lower()
    # Et surtout : l'absence de constat ne vaut pas approbation.
    assert "ne vaut pas approbation" in rest


def test_an_ordinary_draft_keeps_its_own_warning():
    from app.services.export_service import DRAFT_WARNING

    assert _notice(exportable=False, interrupted=False) == DRAFT_WARNING


def test_an_interrupted_run_is_never_presented_as_validated():
    """Même interrompue, une chaîne ne peut pas produire un dossier « contrôlé »."""
    from app.services.export_service import VALIDATED_NOTICE

    assert _notice(exportable=False, interrupted=True) != VALIDATED_NOTICE
    assert _notice(exportable=True, interrupted=False) == VALIDATED_NOTICE


def test_both_formats_use_the_same_selector():
    """Le PDF et le Word ne doivent pas diverger sur ce qui compte le plus."""
    import inspect

    from app.services.export_service import ExportService

    pdf = inspect.getsource(ExportService._dossier_banner)
    docx = inspect.getsource(ExportService._docx_banner)
    assert "_dossier_notice" in pdf
    assert "_dossier_notice" in docx
