"""Tests du budget prévisionnel, du plan de financement et du calendrier.

Trois règles portent le module et sont testées en priorité : aucun montant
n'est deviné, les totaux sont toujours recalculés et jamais saisis, et
« acquis » veut dire acquis.
"""

from __future__ import annotations

import io
import zipfile

import pytest

from app.models.enums import BudgetCategory, ProjectType
from app.services.budget_templates import CATEGORY_ORDER, template_for


def _budget(client, headers, project_id) -> dict:
    response = client.get(f"/api/v1/projects/{project_id}/budget", headers=headers)
    assert response.status_code == 200, response.text
    return response.json()


def _generate(client, headers, project_id, replace: bool = False) -> dict:
    response = client.post(
        f"/api/v1/projects/{project_id}/budget/generate",
        json={"replace": replace},
        headers=headers,
    )
    assert response.status_code == 201, response.text
    return response.json()


def _add_item(client, headers, project_id, **fields) -> dict:
    body = {"category": "PRODUCTION", "label": "Poste", "quantity": 1, "unit_price": 0}
    body.update(fields)
    response = client.post(
        f"/api/v1/projects/{project_id}/budget/items", json=body, headers=headers
    )
    assert response.status_code == 201, response.text
    return response.json()


# ---------------------------------------------------------------------------
# Trame : la structure, jamais les montants
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("project_type", list(ProjectType))
def test_every_project_type_has_a_template_covering_all_phases(project_type):
    lines = template_for(project_type)
    assert lines
    assert {line.category for line in lines} == set(CATEGORY_ORDER)
    # La trame est ordonnée comme les phases de production.
    positions = [CATEGORY_ORDER.index(line.category) for line in lines]
    assert positions == sorted(positions)


def test_the_template_never_carries_an_amount():
    """Un tarif inventé serait un chiffre faux dans un dossier de financement."""
    for project_type in ProjectType:
        for line in template_for(project_type):
            assert not hasattr(line, "unit_price")
            assert line.label and line.label.strip()


def test_generating_installs_the_posts_without_pricing_them(client, auth_headers, project):
    budget = _generate(client, auth_headers, project["id"])

    assert budget["items"], "la trame doit proposer des postes"
    assert budget["total_amount"] == 0.0
    assert all(item["unit_price"] == 0.0 for item in budget["items"])
    assert all(item["amount"] == 0.0 for item in budget["items"])
    # Toutes les phases sont représentées, dans l'ordre de production.
    categories = [total["category"] for total in budget["totals_by_category"]]
    assert categories == [str(c) for c in CATEGORY_ORDER if str(c) in categories]


def test_regenerating_does_not_erase_the_work_already_done(client, auth_headers, project):
    _generate(client, auth_headers, project["id"])
    budget = _budget(client, auth_headers, project["id"])
    first = budget["items"][0]

    client.put(
        f"/api/v1/projects/{project['id']}/budget/items/{first['id']}",
        json={"quantity": 3, "unit_price": 250000},
        headers=auth_headers,
    )
    regenerated = _generate(client, auth_headers, project["id"])

    kept = next(item for item in regenerated["items"] if item["id"] == first["id"])
    assert kept["unit_price"] == 250000
    # Aucun doublon : la trame ne réinstalle pas un poste déjà présent.
    labels = [(item["category"], item["label"]) for item in regenerated["items"]]
    assert len(labels) == len(set(labels))


def test_replacing_starts_from_a_clean_slate(client, auth_headers, project):
    _generate(client, auth_headers, project["id"])
    _add_item(client, auth_headers, project["id"], label="Poste maison", unit_price=1000)

    replaced = _generate(client, auth_headers, project["id"], replace=True)
    assert "Poste maison" not in [item["label"] for item in replaced["items"]]
    assert replaced["total_amount"] == 0.0


# ---------------------------------------------------------------------------
# Totaux : calculés, jamais saisis
# ---------------------------------------------------------------------------
def test_an_item_amount_is_quantity_times_unit_price(client, auth_headers, project):
    item = _add_item(
        client, auth_headers, project["id"], label="Image", quantity=20, unit_price=150000
    )
    assert item["amount"] == 20 * 150000


def test_the_total_follows_the_items(client, auth_headers, project):
    _add_item(client, auth_headers, project["id"], quantity=2, unit_price=100)
    _add_item(client, auth_headers, project["id"], quantity=3, unit_price=1000)

    budget = _budget(client, auth_headers, project["id"])
    assert budget["total_amount"] == 2 * 100 + 3 * 1000


def test_changing_a_price_updates_the_total(client, auth_headers, project):
    item = _add_item(client, auth_headers, project["id"], quantity=1, unit_price=500)
    assert _budget(client, auth_headers, project["id"])["total_amount"] == 500

    client.put(
        f"/api/v1/projects/{project['id']}/budget/items/{item['id']}",
        json={"unit_price": 900},
        headers=auth_headers,
    )
    assert _budget(client, auth_headers, project["id"])["total_amount"] == 900


def test_deleting_an_item_updates_the_total(client, auth_headers, project):
    kept = _add_item(client, auth_headers, project["id"], quantity=1, unit_price=400)
    removed = _add_item(client, auth_headers, project["id"], quantity=1, unit_price=600)
    assert _budget(client, auth_headers, project["id"])["total_amount"] == 1000

    response = client.delete(
        f"/api/v1/projects/{project['id']}/budget/items/{removed['id']}", headers=auth_headers
    )
    assert response.status_code == 200
    budget = _budget(client, auth_headers, project["id"])
    assert budget["total_amount"] == 400
    assert [item["id"] for item in budget["items"]] == [kept["id"]]


def test_category_subtotals_and_shares_add_up(client, auth_headers, project):
    _add_item(client, auth_headers, project["id"], category="PRODUCTION", quantity=1, unit_price=750)
    _add_item(client, auth_headers, project["id"], category="DEVELOPMENT", quantity=1, unit_price=250)

    budget = _budget(client, auth_headers, project["id"])
    totals = {total["category"]: total for total in budget["totals_by_category"]}
    assert totals["PRODUCTION"]["amount"] == 750
    assert totals["DEVELOPMENT"]["amount"] == 250
    assert totals["PRODUCTION"]["share"] == 75.0
    assert sum(total["amount"] for total in budget["totals_by_category"]) == budget["total_amount"]


def test_a_negative_price_is_refused(client, auth_headers, project):
    response = client.post(
        f"/api/v1/projects/{project['id']}/budget/items",
        json={"category": "PRODUCTION", "label": "Poste", "unit_price": -100},
        headers=auth_headers,
    )
    assert response.status_code == 422


# ---------------------------------------------------------------------------
# Plan de financement : « acquis » veut dire acquis
# ---------------------------------------------------------------------------
def _add_source(client, headers, project_id, **fields) -> dict:
    body = {"source_type": "PUBLIC_FUND", "amount": 0, "is_secured": False}
    body.update(fields)
    response = client.post(
        f"/api/v1/projects/{project_id}/funding-plan/lines", json=body, headers=headers
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_the_plan_follows_the_budget_without_being_typed_in(client, auth_headers, project):
    _add_item(client, auth_headers, project["id"], quantity=1, unit_price=1_000_000)

    plan = client.get(
        f"/api/v1/projects/{project['id']}/funding-plan", headers=auth_headers
    ).json()
    assert plan["total_budget"] == 1_000_000


def test_only_secured_sources_count_as_funded(client, auth_headers, project):
    _add_item(client, auth_headers, project["id"], quantity=1, unit_price=1_000_000)
    _add_source(client, auth_headers, project["id"], amount=400_000, is_secured=True,
                source_name="Fonds national")
    _add_source(client, auth_headers, project["id"], amount=300_000, is_secured=False,
                source_name="Coproducteur pressenti")

    plan = client.get(
        f"/api/v1/projects/{project['id']}/funding-plan", headers=auth_headers
    ).json()

    assert plan["secured_amount"] == 400_000
    # L'espéré est compté à part : il ne gonfle pas le financement acquis.
    assert plan["identified_amount"] == 700_000
    assert plan["sought_amount"] == 600_000
    assert plan["funded_percentage"] == 40.0
    assert plan["uncovered_amount"] == 300_000


def test_securing_a_source_moves_the_coverage(client, auth_headers, project):
    _add_item(client, auth_headers, project["id"], quantity=1, unit_price=1000)
    line = _add_source(client, auth_headers, project["id"], amount=1000, is_secured=False)

    client.put(
        f"/api/v1/projects/{project['id']}/funding-plan/lines/{line['id']}",
        json={"is_secured": True},
        headers=auth_headers,
    )
    plan = client.get(
        f"/api/v1/projects/{project['id']}/funding-plan", headers=auth_headers
    ).json()
    assert plan["funded_percentage"] == 100.0
    assert plan["sought_amount"] == 0.0


def test_an_empty_budget_reports_zero_percent_rather_than_dividing_by_zero(
    client, auth_headers, project
):
    plan = client.get(
        f"/api/v1/projects/{project['id']}/funding-plan", headers=auth_headers
    ).json()
    assert plan["total_budget"] == 0.0
    assert plan["funded_percentage"] == 0.0


# ---------------------------------------------------------------------------
# Calendrier
# ---------------------------------------------------------------------------
def test_a_phase_is_updated_not_duplicated(client, auth_headers, project):
    for dates in (("2027-01-01", "2027-02-01"), ("2027-03-01", "2027-04-01")):
        response = client.put(
            f"/api/v1/projects/{project['id']}/schedule",
            json={"phase": "PRODUCTION", "start_date": dates[0], "end_date": dates[1]},
            headers=auth_headers,
        )
        assert response.status_code == 200, response.text

    schedule = client.get(
        f"/api/v1/projects/{project['id']}/schedule", headers=auth_headers
    ).json()
    assert len(schedule) == 1
    assert schedule[0]["start_date"] == "2027-03-01"


def test_the_schedule_follows_the_order_of_production(client, auth_headers, project):
    for phase in ("DISTRIBUTION", "DEVELOPMENT", "PRODUCTION"):
        client.put(
            f"/api/v1/projects/{project['id']}/schedule",
            json={"phase": phase},
            headers=auth_headers,
        )
    schedule = client.get(
        f"/api/v1/projects/{project['id']}/schedule", headers=auth_headers
    ).json()
    assert [row["phase"] for row in schedule] == ["DEVELOPMENT", "PRODUCTION", "DISTRIBUTION"]


# ---------------------------------------------------------------------------
# Isolation et score
# ---------------------------------------------------------------------------
def test_a_budget_belongs_to_its_project_owner_alone(client, make_user):
    alice, _ = make_user("alice.budget@example.com", plan="PRO_AUTHOR")
    bob, _ = make_user("bob.budget@example.com", plan="PRO_AUTHOR")
    project_id = client.post(
        "/api/v1/projects",
        json={"title": "Projet d'Alice", "project_type": "DOCUMENTARY"},
        headers=alice,
    ).json()["id"]

    # 404 et non 403 : la réponse ne révèle pas que le projet existe.
    assert client.get(f"/api/v1/projects/{project_id}/budget", headers=bob).status_code == 404
    assert (
        client.post(
            f"/api/v1/projects/{project_id}/budget/items",
            json={"category": "PRODUCTION", "label": "Intrusion"},
            headers=bob,
        ).status_code
        == 404
    )


def test_chiffrer_le_budget_fait_monter_le_score_de_maturite(client, auth_headers, project):
    before = client.get(
        f"/api/v1/projects/{project['id']}/score", headers=auth_headers
    ).json()

    _generate(client, auth_headers, project["id"])
    budget = _budget(client, auth_headers, project["id"])
    for item in budget["items"][:6]:
        client.put(
            f"/api/v1/projects/{project['id']}/budget/items/{item['id']}",
            json={"unit_price": 100000},
            headers=auth_headers,
        )

    after = client.get(f"/api/v1/projects/{project['id']}/score", headers=auth_headers).json()
    assert after["total"] > before["total"]


# ---------------------------------------------------------------------------
# Export tableur
# ---------------------------------------------------------------------------
def test_the_xlsx_export_is_reserved_to_plans_that_include_it(client, make_user):
    headers, _ = make_user("gratuit.budget@example.com")  # plan FREE : allows_export=False
    project_id = client.post(
        "/api/v1/projects",
        json={"title": "Projet test", "project_type": "DOCUMENTARY"},
        headers=headers,
    ).json()["id"]

    response = client.get(f"/api/v1/projects/{project_id}/export/xlsx", headers=headers)
    assert response.status_code == 402
    assert response.json()["code"] == "export_not_included"


def test_the_xlsx_export_carries_the_three_sheets_and_live_formulas(
    client, auth_headers, project
):
    _add_item(
        client, auth_headers, project["id"], category="PRODUCTION", label="Image",
        quantity=20, unit_price=150000,
    )
    _add_source(client, auth_headers, project["id"], amount=1_000_000, is_secured=True,
                source_name="Fonds national")
    client.put(
        f"/api/v1/projects/{project['id']}/schedule",
        json={"phase": "PRODUCTION", "start_date": "2027-05-01"},
        headers=auth_headers,
    )

    response = client.get(f"/api/v1/projects/{project['id']}/export/xlsx", headers=auth_headers)
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Budget", "Plan de financement", "Calendrier"]

    # Les montants sont des formules : un prix corrigé dans le tableur met le
    # total à jour, au lieu de laisser un chiffre figé qui ne correspond plus.
    formulas = [
        cell.value
        for row in workbook["Budget"].iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    assert any(formula.startswith("=C") and "*E" in formula for formula in formulas)
    assert any(formula.startswith("=SUM(") for formula in formulas)

    plan_values = [
        cell.value for row in workbook["Plan de financement"].iter_rows() for cell in row
    ]
    assert "Fonds national" in plan_values
    assert "Financement acquis" in plan_values


def test_an_xlsx_file_is_a_readable_archive(client, auth_headers, project):
    """Garde-fou de format : un .xlsx est une archive ZIP contenant du XML."""
    response = client.get(f"/api/v1/projects/{project['id']}/export/xlsx", headers=auth_headers)
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert "xl/workbook.xml" in archive.namelist()


def test_the_export_says_what_to_do_when_the_budget_is_empty(client, auth_headers, project):
    from openpyxl import load_workbook

    response = client.get(f"/api/v1/projects/{project['id']}/export/xlsx", headers=auth_headers)
    workbook = load_workbook(io.BytesIO(response.content))
    values = [cell.value for row in workbook["Budget"].iter_rows() for cell in row]
    assert any(isinstance(value, str) and "Budget vide" in value for value in values)


def test_budget_categories_cover_the_whole_production(client, auth_headers, project):
    """Garde-fou : ajouter une phase au modèle sans l'ordonner casserait l'affichage."""
    assert set(CATEGORY_ORDER) == set(BudgetCategory)
